"""
Cliente DataWeb assíncrono — substitui chamada síncrona a /api/v1/processar.

Deploy no Scraper (212.47.68.222):
  cp infra/scraper-integration/dataweb_client.py /opt/apps/scraper/repo/app/services/dataweb_client.py
  systemctl restart scraper   # ou docker compose restart, conforme o deploy do scraper

Variáveis de ambiente (opcional):
  DATAWEB_BASE_URL=http://127.0.0.1:5267/dataweb
  DATAWEB_POLL_SECONDS=3
  DATAWEB_MAX_WAIT_SECONDS=7200
  DATAWEB_LOTE_MAX=500
"""

from __future__ import annotations

import logging
import os
import time
from typing import Any

import requests

logger = logging.getLogger(__name__)


class DataWebError(Exception):
    """Erro retornado pelo DataWeb ou falha de comunicação."""


class DataWebClient:
    def __init__(
        self,
        base_url: str | None = None,
        poll_seconds: float | None = None,
        max_wait_seconds: int | None = None,
        lote_max: int | None = None,
    ) -> None:
        self.base_url = (base_url or os.environ.get("DATAWEB_BASE_URL", "http://127.0.0.1:5267/dataweb")).rstrip("/")
        self.poll_seconds = float(poll_seconds or os.environ.get("DATAWEB_POLL_SECONDS", "3"))
        self.max_wait_seconds = int(max_wait_seconds or os.environ.get("DATAWEB_MAX_WAIT_SECONDS", "7200"))
        self.lote_max = int(lote_max or os.environ.get("DATAWEB_LOTE_MAX", "500"))
        self._session = requests.Session()

    def processar_cnjs(self, cnjs: list[str]) -> bytes:
        """Processa uma lista de CNJs e devolve bytes do Excel (fluxo assíncrono)."""
        cnjs = [c.strip() for c in cnjs if c and str(c).strip()]
        if not cnjs:
            raise DataWebError("Nenhum CNJ informado.")

        if len(cnjs) <= self.lote_max:
            return self._processar_lote(cnjs)

        lotes = [cnjs[i : i + self.lote_max] for i in range(0, len(cnjs), self.lote_max)]
        logger.info("DataWeb: %s CNJ(s) em %s lote(s) (assíncrono)", len(cnjs), len(lotes))

        partes: list[bytes] = []
        for idx, lote in enumerate(lotes, start=1):
            logger.info("DataWeb: lote %s/%s (%s CNJs)", idx, len(lotes), len(lote))
            partes.append(self._processar_lote(lote))

        if len(partes) == 1:
            return partes[0]

        return _mesclar_excels(partes)

    def _processar_lote(self, cnjs: list[str]) -> bytes:
        job_id = self._criar_job(cnjs)
        self._aguardar_job(job_id)
        return self._baixar_excel(job_id)

    def _criar_job(self, cnjs: list[str]) -> str:
        url = f"{self.base_url}/api/v1/processar/json"
        try:
            resp = self._session.post(url, json={"cnjs": cnjs}, timeout=60)
        except requests.RequestException as exc:
            raise DataWebError(f"Falha ao criar job no DataWeb: {exc}") from exc

        if resp.status_code not in (200, 202):
            raise DataWebError(f"Erro HTTP {resp.status_code} ao criar job: {resp.text[:500]}")

        data = resp.json()
        job_id = data.get("jobId") or data.get("JobId")
        if not job_id:
            raise DataWebError(f"Resposta sem jobId: {data}")
        return str(job_id)

    def _aguardar_job(self, job_id: str) -> dict[str, Any]:
        url = f"{self.base_url}/api/v1/processar/json/{job_id}/status"
        deadline = time.time() + self.max_wait_seconds

        while time.time() < deadline:
            try:
                resp = self._session.get(url, timeout=30)
            except requests.RequestException as exc:
                raise DataWebError(f"Falha ao consultar status do job {job_id}: {exc}") from exc

            if resp.status_code == 404:
                raise DataWebError(f"Job não encontrado: {job_id}")

            if resp.status_code >= 400:
                raise DataWebError(f"Erro HTTP {resp.status_code} no status: {resp.text[:500]}")

            status = resp.json()
            st = (status.get("status") or status.get("Status") or "").lower()
            cnjs_proc = status.get("cnjsProcessados") or status.get("CnjsProcessados") or 0
            total = status.get("totalCnjs") or status.get("TotalCnjs") or "?"
            prog = status.get("progresso") or status.get("Progresso") or 0
            logger.info("DataWeb job %s: %s — %s/%s CNJs (%s%%)", job_id, st, cnjs_proc, total, prog)

            if st == "concluido":
                return status

            if st == "erro":
                erro = status.get("erro") or status.get("Erro") or "Erro desconhecido"
                raise DataWebError(f"Job {job_id} falhou: {erro}")

            time.sleep(self.poll_seconds)

        raise DataWebError(
            f"Job {job_id} não concluiu em {self.max_wait_seconds}s. "
            "Aumente DATAWEB_MAX_WAIT_SECONDS se necessário."
        )

    def _baixar_excel(self, job_id: str) -> bytes:
        url = f"{self.base_url}/api/v1/processar/json/{job_id}/excel"
        try:
            resp = self._session.get(url, timeout=300)
        except requests.RequestException as exc:
            raise DataWebError(f"Falha ao baixar Excel do job {job_id}: {exc}") from exc

        if resp.status_code == 409:
            raise DataWebError(f"Job {job_id} ainda não concluído ao baixar Excel.")

        if resp.status_code >= 400:
            raise DataWebError(f"Erro HTTP {resp.status_code} ao baixar Excel: {resp.text[:500]}")

        if not resp.content:
            raise DataWebError(f"Excel vazio para job {job_id}.")

        return resp.content

    # Aliases para compatibilidade com rotas antigas do Scraper
    def processar_excel(self, cnjs: list[str]) -> bytes:
        return self.processar_cnjs(cnjs)

    def consultar(self, cnjs: list[str]) -> bytes:
        return self.processar_cnjs(cnjs)


def processar(cnjs: list[str], base_url: str | None = None) -> bytes:
    """Atalho module-level (compatível com imports antigos)."""
    return DataWebClient(base_url=base_url).processar_cnjs(cnjs)


def _mesclar_excels(partes: list[bytes]) -> bytes:
    """Mescla vários .xlsx (mesma estrutura DataWeb) em um único arquivo."""
    try:
        from io import BytesIO

        from openpyxl import load_workbook
    except ImportError as exc:
        raise DataWebError(
            "Para mesclar lotes, instale openpyxl no Scraper: pip install openpyxl"
        ) from exc

    wb_dest = load_workbook(BytesIO(partes[0]))
    ws_dest = wb_dest.active
    if ws_dest is None:
        raise DataWebError("Planilha DataWeb sem aba ativa.")

    for blob in partes[1:]:
        wb_src = load_workbook(BytesIO(blob), read_only=True, data_only=True)
        ws_src = wb_src.active
        if ws_src is not None:
            for row in ws_src.iter_rows(min_row=2, values_only=True):
                if row and any(cell is not None and str(cell).strip() for cell in row):
                    ws_dest.append(list(row))
        wb_src.close()

    out = BytesIO()
    wb_dest.save(out)
    return out.getvalue()
